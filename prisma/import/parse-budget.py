#!/usr/bin/env python3
"""Parse a budget workbook into prisma/import/budget-2026.json.

    python3 prisma/import/parse-budget.py [workbook.xlsx]

Each committee sheet (AKTU01, SEX02, …) becomes a cost center; its Konto rows
become accounts and the indented rows beneath become line items. Two value
columns are read per sheet: J — the budget being written, and the only one
imported (see import-budget.ts) — and L, the one it was revised from, kept here
for reference. Both name themselves: Parametrar!I5 and I6 hold the two budgets'
names ("2026 DELTA", "2026 VTM"), and every sheet's column heading points at
them, so the JSON carries the names rather than this file guessing at them.

Where a line uses the "antal × á-pris" form (cells F × H) the imported column
keeps that structure, mapping well-known Parametrar references to named
variables. Anything else — and the whole previous column — is stored as its
computed value, so the totals always match the spreadsheet exactly.
"""
import json
import re
import os
import sys
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(__file__)
DEFAULT_XLSX = os.path.join(HERE, "..", "..", "VTM-EXTRA.xlsx")
XLSX = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
OUT = os.path.join(HERE, "budget-2026.json")

# Parametrar rows we expose as named variables, keyed by the label in column A.
# Located by label and not by row number: each revision of the workbook inserts
# rows into this sheet, and a row number that has drifted does not fail — it
# names the wrong cell, and the budget imports plausible numbers under the wrong
# variable. A label that no longer matches stops the parse instead (see
# resolve_params).
PARAM_BY_LABEL = {
    "Prisbasbelopp": "PRISBASBELOPP",
    "Tack enl. policy": "TACK_POLICY",
    "Tröja enl. policy": "TROJA_POLICY",
    "Mat": "MAT",
    "Lokalbokning för lunchföreläsning": "LOKAL_LUNCHFORELASNING",
    "Fortnox — Bokföring": "FORTNOX_BOKFORING",
    "Fortnox — Bokföring (för bokslut)": "FORTNOX_BOKFORING_BOKSLUT",
    "Fortnox — Offert och order": "FORTNOX_OFFERT_ORDER",
    "Fortnox — Fakturering": "FORTNOX_FAKTURERING",
    "Fortnox — Läs (Revisorer)": "FORTNOX_LAS",
    "Fortnox — Integration": "FORTNOX_INTEGRATION",
    "GENOMSNITTLIGT Påslag öl": "PASLAG_OL",
    "GENOMSNITTLIGT Påslag övrig jäst dryck": "PASLAG_OVRIG_JAST",
    "GENOMSNITTLIGT Påslag vin": "PASLAG_VIN",
    "GENOMSNITTLIGT Påslag spritdryck": "PASLAG_SPRIT",
}

# Filled in by resolve_params(): {row: variable name}.
PARAM_BY_ROW = {}

# Rows the sheets reference that no label above claims. Their value still lands
# in the budget (as a number rather than a live variable reference), so this is
# a report, not an error — see the note printed by main().
UNNAMED_PARAMS = []

COL_F, COL_H, COL_J, COL_L = 6, 8, 10, 12
SKIP_B = {"Summa", "Intäkter", "Kostnader", "Resultat", ""}

# Sexmästeriet's sheets book alcohol under a *range* header — "3011-3014",
# "4011-4014" — with one line per drink beneath it. BAS numbers those
# individually and so does the ledger, so a range is split back into one account
# per line, matched on the line's own description (which is the account's name
# in the chart of accounts). Kept as a range, the code is not a number: the
# Utfall report resolves an account to a section with Number(code), so every
# such row was silently dropped from the budget comparison.
RANGE_HEADER = re.compile(r"^(\d{4})\s*-\s*(\d{4})$")

# Offset from the range's first account, by line description (lowercased).
DRINK_OFFSET = {
    "öl": 0,
    "övrig jäst dryck": 1,
    "vin": 2,
    "spritdryck": 3,
}

# Descriptions under a range header that no rule could place. Reported at the
# end rather than guessed at — a drink on the wrong account is a silent error.
UNPLACED = []


def split_range_account(first, last, description):
    """The real account a line under a range header belongs to, or None."""
    offset = DRINK_OFFSET.get(description.strip().lower())
    if offset is None:
        return None
    code = first + offset
    return str(code) if code <= last else None

wb = openpyxl.load_workbook(XLSX, data_only=False)
wbv = openpyxl.load_workbook(XLSX, data_only=True)


def resolve_params():
    """Find each named parameter's row by its label in Parametrar column A.

    Every label must match exactly one row: a missing one means the workbook
    renamed it (and the lines referencing it would quietly become plain numbers),
    a repeated one means the label no longer identifies a single parameter.
    Either way the mapping is no longer trustworthy, so stop.
    """
    ws = wbv["Parametrar"]
    rows_by_label = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip() in PARAM_BY_LABEL:
            rows_by_label.setdefault(a.strip(), []).append(r)

    problems = []
    for label, name in PARAM_BY_LABEL.items():
        rows = rows_by_label.get(label, [])
        if len(rows) != 1:
            problems.append(f"    {name:26} {label!r}: "
                            f"{'no such row' if not rows else f'{len(rows)} rows {rows}'}")
        else:
            PARAM_BY_ROW[rows[0]] = name
    if problems:
        raise SystemExit(
            "Parametrar no longer matches PARAM_BY_LABEL in this parser:\n"
            + "\n".join(problems)
            + "\n\nFix the labels above to match the workbook and re-run."
        )


def collect_unnamed_params():
    """Parametrar cells the sheets reference that aren't named variables.

    Informational: an unnamed reference is imported as the number it evaluates
    to, which is correct but dead — changing the parameter won't reflow it.
    """
    ws = wbv["Parametrar"]
    seen = set()
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, ArrayFormula):
                    v = v.text
                if not isinstance(v, str) or "Parametrar" not in v:
                    continue
                for m in re.finditer(r"Parametrar!\$?I\$?(\d+)", v):
                    r = int(m.group(1))
                    if r in PARAM_BY_ROW or r in seen or r <= 6:
                        continue  # rows 1–6 are the sheet's own headings
                    seen.add(r)
                    UNNAMED_PARAMS.append((r, str(ws.cell(r, 1).value or "").strip()))
    UNNAMED_PARAMS.sort()


def budget_names():
    """The two columns' names, as the workbook states them (Parametrar I5/I6)."""
    ws = wbv["Parametrar"]
    current = str(ws.cell(5, 9).value or "").strip()
    previous = str(ws.cell(6, 9).value or "").strip()
    if not current:
        raise SystemExit("Parametrar!I5 (Nuvarande budgetnamn) is empty — "
                         "the imported budget would have no name.")
    return current, previous


def year_of(name):
    """The year a budget name opens with ("2026 DELTA" → 2026)."""
    m = re.match(r"\s*(\d{4})\b", name)
    if not m:
        raise SystemExit(f"Budget name {name!r} does not start with a year, so "
                         f"there is nothing to file the budget under.")
    return int(m.group(1))


def fmtnum(v):
    """Format a number as a clean string ("4002", "2", "1.44", "560705.0151")."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v).strip()
    if f == int(f):
        return str(int(f))
    return repr(f)


def numeric(v):
    """A cell's value as a float, or None if it isn't a plain number."""
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def param_ref(formula):
    """Map '=Parametrar!$I$8' → 'TACK_POLICY', else None."""
    if not isinstance(formula, str):
        return None
    m = re.match(r"^=\s*Parametrar!\$?I\$?(\d+)\s*$", formula)
    if not m:
        return None
    return PARAM_BY_ROW.get(int(m.group(1)))


def cell_term(sheet_f, sheet_v, r, c):
    """Resolve a quantity/á-pris cell to an expression string.

    A Parametrar reference becomes its variable name; anything else becomes the
    computed numeric value. Returns None when the cell is empty.
    """
    raw = sheet_f.cell(r, c).value
    if raw is None or raw == "":
        return None
    if isinstance(raw, str) and raw.startswith("="):
        var = param_ref(raw)
        if var:
            return var
        return fmtnum(sheet_v.cell(r, c).value)
    if isinstance(raw, ArrayFormula):
        return fmtnum(sheet_v.cell(r, c).value)
    return fmtnum(raw)


def value_expr(sheet_f, sheet_v, r, c):
    """A standalone amount cell → a formula referencing the named variable if
    it's a plain Parametrar ref, otherwise the computed value as a string.

    The "=" matters: an amount is only evaluated as a formula when it starts
    with one (see lib/budget/eval.ts), so a bare "TACK_POLICY" is read as a
    number, fails, and the line silently counts as zero.
    """
    raw = sheet_f.cell(r, c).value
    if isinstance(raw, str) and raw.startswith("="):
        var = param_ref(raw)
        if var:
            return f"={var}"
    val = sheet_v.cell(r, c).value
    return fmtnum(val) if val is not None else "0"


def result_unit_map():
    """Build {sheetCode: resultUnitName} from the Rambudget overview.

    The overview groups the cost-centre sheets under a heading per
    resultatenhet — Sexmästeriet, Källarmästeriet — so the heading a sheet sits
    under is the unit it belongs to."""
    ws = wb["Rambudget"]
    sheet_names = set(wb.sheetnames)
    out, current = {}, None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if not isinstance(a, str) or not a.strip():
            continue
        a = a.strip()
        if a in sheet_names:
            if current:
                out[a] = current
        elif a.startswith("Totalt") or a in ("Beteckning", "Allmän info", "Totalt"):
            continue
        else:
            current = a  # a result-unit header row
    return out


def parse_variables():
    ws = wbv["Parametrar"]
    out = []
    for row, name in PARAM_BY_ROW.items():
        out.append({
            "name": name,
            "rev": fmtnum(ws.cell(row, COL_J - 1).value) or "0",   # column I
            "orig": fmtnum(ws.cell(row, COL_L - 1).value) or "0",  # column K
        })
    return out


def parse_sheet(code):
    sf, sv = wb[code], wbv[code]
    name = sf.cell(1, 2).value or code
    # find the "Konto" header row
    start = None
    for r in range(1, sf.max_row + 1):
        if str(sf.cell(r, 1).value).strip() == "Konto":
            start = r + 1
            break
    if start is None:
        start = 11

    accounts, by_code, current = [], {}, None
    for r in range(start, sf.max_row + 1):
        a = sf.cell(r, 1).value
        b = sf.cell(r, 2).value
        b_str = str(b).strip() if b is not None else ""

        if a is not None and str(a).strip() != "" and b_str not in ("Intäkter", "Kostnader", "Resultat"):
            # account header
            acode = fmtnum(a) if isinstance(a, (int, float)) else str(a).strip()
            rng = RANGE_HEADER.match(acode)
            if rng:
                # Not an account of its own: each line beneath gets its own.
                current = {"range": (int(rng.group(1)), int(rng.group(2)))}
                continue
            if acode in by_code:
                current = by_code[acode]
            else:
                current = {"code": acode, "name": b_str, "lineItems": []}
                by_code[acode] = current
                accounts.append(current)
            continue

        if b_str in SKIP_B or current is None:
            continue

        # Under a range header the account is the line's own, named for it.
        if "range" in current:
            first, last = current["range"]
            acode = split_range_account(first, last, b_str)
            if acode is None:
                UNPLACED.append((code, f"{first}-{last}", b_str))
                continue
            if acode in by_code:
                target = by_code[acode]
            else:
                target = {"code": acode, "name": b_str, "lineItems": []}
                by_code[acode] = target
                accounts.append(target)
        else:
            target = current

        # line item. The amount is always column J (REV) / L (baseline) — these
        # are the sheet's source of truth (a row may override the formula with a
        # literal). The "antal × á-pris" form is kept for display only when the
        # row is a clean 2-factor F × H whose product matches J (no extra D
        # factor, no manual override), so the breakdown can never disagree with
        # the total.
        fv = numeric(sv.cell(r, COL_F).value)
        hv = numeric(sv.cell(r, COL_H).value)
        jv = numeric(sv.cell(r, COL_J).value)
        d_empty = sf.cell(r, 4).value in (None, "")
        structured = (fv is not None and hv is not None and d_empty
                      and jv is not None and abs(fv * hv - jv) < 1.0)
        if structured:
            rev = {"quantity": cell_term(sf, sv, r, COL_F),
                   "unitPrice": cell_term(sf, sv, r, COL_H),
                   "expression": value_expr(sf, sv, r, COL_J)}
        else:
            rev = {"quantity": None, "unitPrice": None, "expression": value_expr(sf, sv, r, COL_J)}
        orig = {"quantity": None, "unitPrice": None, "expression": value_expr(sf, sv, r, COL_L)}
        target["lineItems"].append({"description": b_str, "rev": rev, "orig": orig})

    # drop empty accounts (header with no lines and no name carry no signal)
    accounts = [a for a in accounts if a["lineItems"]]
    return {"code": code, "name": str(name).strip(), "lineItems_unused": False, "accounts": accounts}


def rambudget_results():
    """Each sheet's Resultat as the Rambudget overview states it: {code: {rev, orig}}."""
    ws = wbv["Rambudget"]
    names = set(wb.sheetnames)
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip() in names:
            out[a.strip()] = {"rev": ws.cell(r, 6).value, "orig": ws.cell(r, 10).value}
    return out, {"rev": ws.cell(3, 6).value, "orig": ws.cell(3, 10).value}


def eval_expr(expr, variables):
    """What the app will make of a stored expression (see lib/budget/eval.ts)."""
    s = (expr or "").strip()
    if not s:
        return 0.0
    if s.startswith("="):
        s = s[1:].strip()
        if s in variables:
            return eval_expr(variables[s], variables)
        if "*" in s:
            out = 1.0
            for part in s.split("*"):
                out *= eval_expr(part.strip(), variables)
            return out
    try:
        return float(s)
    except ValueError:
        return float("nan")


def reconcile(data):
    """Check the parsed budget against the spreadsheet's own totals.

    The point of the exercise: a sheet whose line items don't add up to the
    Resultat it reports is either a parser that lost rows or a spreadsheet that
    was summarised by hand, and both are worth stopping for. Silently importing
    the difference is how a budget comparison ends up quietly wrong.
    """
    ram, ram_total = rambudget_results()
    ok = True
    for side in ("rev", "orig"):
        variables = {v["name"]: v[side] for v in data["variables"]}
        bad, total = [], 0.0
        for cc in data["costCenters"]:
            got = 0.0
            for a in cc["accounts"]:
                sign = 1 if a["code"].strip().startswith("3") else -1
                for li in a["lineItems"]:
                    got += sign * eval_expr(li[side]["expression"], variables)
            total += got
            want = ram.get(cc["code"], {}).get(side)
            if want is None or abs(got - want) > 1:
                bad.append((cc["code"], got, want))
        # Only the left column is imported (see import-budget.ts), so only it
        # can make this fail; the previous column is reported for information.
        imported = side == "rev"
        label = data["name"] if imported else f"{data['previousName']} (importeras ej)"
        stated = ram_total[side] or 0
        reconciled = not bad and abs(total - stated) <= 1
        mark = "OK " if reconciled else ("!! " if imported else "-- ")
        print(f"{mark}{label:30} parsed {total:>14,.2f}   rambudget {stated:>14,.2f}   "
              f"diff {total - stated:>12,.2f}")
        for code, got, want in bad:
            if imported:
                ok = False
            print(f"     {code:8} parsed {got:>13,.2f}   sheet states {want if want is None else f'{want:,.2f}':>13}"
                  f"   diff {got - (want or 0):>12,.2f}")
    return ok


def main():
    print(f"Reading {os.path.relpath(XLSX)}")
    resolve_params()
    collect_unnamed_params()
    name, previous_name = budget_names()

    units = result_unit_map()
    skip = {"Rambudget", "Parametrar"}
    cost_centers = []
    for code in wb.sheetnames:
        if code in skip:
            continue
        cc = parse_sheet(code)
        cc["resultUnit"] = units.get(code)
        del cc["lineItems_unused"]
        cost_centers.append(cc)

    data = {
        "year": year_of(name),
        "name": name,
        "previousName": previous_name,
        "variables": parse_variables(),
        "costCenters": cost_centers,
    }
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)

    if UNNAMED_PARAMS:
        print(f"Note: {len(UNNAMED_PARAMS)} referenced Parametrar row(s) have no "
              f"named variable, so lines using them import as fixed numbers:")
        for r, label in UNNAMED_PARAMS:
            print(f"    I{r:<4} {label}")
        print()

    if UNPLACED:
        print(f"WARNING: {len(UNPLACED)} line(s) under a range header could not be "
              f"placed on an account and were dropped:")
        for cc, rng, desc in UNPLACED:
            print(f"    {cc} {rng}: {desc!r}")

    n_acc = sum(len(c["accounts"]) for c in cost_centers)
    n_li = sum(len(a["lineItems"]) for c in cost_centers for a in c["accounts"])
    print(f"Wrote {OUT}: \"{name}\" (revised from \"{previous_name}\") — "
          f"{len(cost_centers)} cost centers, {n_acc} accounts, {n_li} line items, "
          f"{len(data['variables'])} variables\n")

    if not reconcile(data):
        print("\nThe sheets above do not add up to the Resultat they report, so the\n"
              "budget imported from them would be wrong by that much. The JSON was\n"
              "still written; fix the workbook and re-run.")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
