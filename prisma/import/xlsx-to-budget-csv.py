#!/usr/bin/env python3
"""Convert a budget workbook to the interchange CSV that /api/budgets/import reads.

    python3 prisma/import/xlsx-to-budget-csv.py [workbook.xlsx] [-c J|L] [-o out.csv]

This script knows one workbook shape: the VTM/DELTA family, where each
committee sheet (AKTU01, SEX02, …) is a cost centre, column A carries Konto
numbers, column B the names and descriptions, and the amounts sit in column J
(the revision being written) or L (the one it was revised from).

It is deliberately NOT the importer. Older budgets are laid out differently and
each will need its own conversion, or a CSV assembled by hand in Excel, which
for a one-off is the faster route. What survives from year to year is the CSV,
which is why the app reads that and not a workbook: see lib/domain/budget-csv.ts.

Two things are dropped on purpose:

  * Variables. The CSV has no Parametrar column, so a line referencing
    Prisbasbelopp becomes the number it evaluated to. For a past budget that is
    what you want. Last year's budget should not reflow when this year's
    prisbasbelopp is set.
  * The "previous column" pairing. One run emits one budget; run it twice with
    -c J and -c L to get both, and import them under their own names.

Every sheet is reconciled against the Resultat the Rambudget overview states for
it, and a mismatch is reported. A budget that disagrees with the budget that was
voted through is worse than no budget at all, so check the report before
importing.
"""
import argparse
import csv
import os
import re
import sys

import openpyxl

# Column layout of a committee sheet, 1-based (openpyxl's convention).
COL_A, COL_B, COL_D, COL_F, COL_H = 1, 2, 4, 6, 8
COL_J, COL_L = 10, 12

# Column B values that are not line items: the sheets' own subtotal rows.
SKIP_B = {"Summa", "Intäkter", "Kostnader", "Resultat", ""}

# Sheets that are not cost centres.
NON_SHEETS = {"Rambudget", "Parametrar"}

# Sexmästeriet books alcohol under a range header ("3011-3014") with one line
# per drink beneath it. BAS numbers those individually and so does the ledger,
# so the range is split back into one account per line, matched on the line's
# own description. Kept as a range the code is not a number, and the Utfall
# report resolves an account with Number(code), so such a row would be dropped
# from the comparison without a word.
RANGE_HEADER = re.compile(r"^(\d{4})\s*-\s*(\d{4})$")
DRINK_OFFSET = {"öl": 0, "övrig jäst dryck": 1, "vin": 2, "spritdryck": 3}


def fmtnum(v):
    """A number as a clean string: 4002, 2, 1.44, 560705.0151."""
    if v is None or isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v).strip()
    return str(int(f)) if f == int(f) else repr(f)


def numeric(v):
    """A cell's value as a float, or None if it is not a plain number."""
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def split_range_account(first, last, description):
    """The real account a line under a range header belongs to, or None."""
    offset = DRINK_OFFSET.get(description.strip().lower())
    if offset is None:
        return None
    code = first + offset
    return str(code) if code <= last else None


def parse_sheet(wb, wbv, code, amount_col, unplaced):
    """Every line of one committee sheet, as interchange-CSV rows."""
    sf, sv = wb[code], wbv[code]
    cc_name = str(sf.cell(1, COL_B).value or code).strip()

    start = None
    for r in range(1, sf.max_row + 1):
        if str(sf.cell(r, COL_A).value).strip() == "Konto":
            start = r + 1
            break
    if start is None:
        return cc_name, []

    rows = []
    names_by_code = {}   # account code → the name its header gave it
    current = None       # {"code": …} or {"range": (first, last)}

    for r in range(start, sf.max_row + 1):
        a = sf.cell(r, COL_A).value
        b = sf.cell(r, COL_B).value
        b_str = str(b).strip() if b is not None else ""

        # An account header: a code in column A, unless the row is one of the
        # sheet's own summary lines.
        if a is not None and str(a).strip() != "" and b_str not in ("Intäkter", "Kostnader", "Resultat"):
            acode = fmtnum(a) if isinstance(a, (int, float)) else str(a).strip()
            rng = RANGE_HEADER.match(acode or "")
            if rng:
                current = {"range": (int(rng.group(1)), int(rng.group(2)))}
            else:
                current = {"code": acode}
                names_by_code.setdefault(acode, b_str)
            continue

        if b_str in SKIP_B or current is None:
            continue

        if "range" in current:
            first, last = current["range"]
            acode = split_range_account(first, last, b_str)
            if acode is None:
                unplaced.append((code, f"{first}-{last}", b_str))
                continue
            # Under a range header the line's own description names the account.
            names_by_code.setdefault(acode, b_str)
        else:
            acode = current["code"]

        # The amount column is the sheet's source of truth: a row may override
        # its own formula with a literal. The antal x á-pris pair is carried
        # over only when it is a clean two-factor F x H whose product matches
        # that column (no extra D factor, no manual override), so the breakdown
        # can never disagree with the total it is shown beside.
        fv = numeric(sv.cell(r, COL_F).value)
        hv = numeric(sv.cell(r, COL_H).value)
        av = numeric(sv.cell(r, amount_col).value)
        d_empty = sf.cell(r, COL_D).value in (None, "")
        structured = (
            fv is not None and hv is not None and d_empty
            and av is not None and abs(fv * hv - av) < 1.0
        )

        if structured:
            rows.append({
                "account": acode, "account_name": names_by_code.get(acode, acode),
                "description": b_str,
                "quantity": fmtnum(fv), "unit_price": fmtnum(hv), "amount": "",
                "_value": fv * hv,
            })
        else:
            rows.append({
                "account": acode, "account_name": names_by_code.get(acode, acode),
                "description": b_str,
                "quantity": "", "unit_price": "", "amount": fmtnum(av) or "0",
                "_value": av or 0.0,
            })

    # A header whose account never took a line carries no signal; nothing to do,
    # since rows are only emitted for lines in the first place.
    for row in rows:
        row["account_name"] = names_by_code.get(row["account"], row["account"])
    return cc_name, rows


def rambudget_results(wb, wbv, amount_col):
    """Each sheet's Resultat as the Rambudget overview states it: {code: value}."""
    ws = wbv["Rambudget"]
    names = set(wb.sheetnames)
    # The overview reports the two columns in F and J, mirroring J and L on the
    # committee sheets.
    col = 6 if amount_col == COL_J else 10
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, COL_A).value
        if isinstance(a, str) and a.strip() in names:
            out[a.strip()] = numeric(ws.cell(r, col).value)
    return out


def main():
    here = os.path.dirname(__file__)
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", nargs="?",
                    default=os.path.join(here, "..", "..", "VTM-EXTRA.xlsx"))
    ap.add_argument("-c", "--column", choices=["J", "L"], default="J",
                    help="which amount column to read (default J, the revision being written)")
    ap.add_argument("-o", "--out", default=os.path.join(here, "budget.csv"))
    args = ap.parse_args()

    amount_col = COL_J if args.column == "J" else COL_L
    wb = openpyxl.load_workbook(args.workbook, data_only=False)
    wbv = openpyxl.load_workbook(args.workbook, data_only=True)

    unplaced = []
    out_rows = []
    totals = {}
    for code in wb.sheetnames:
        if code in NON_SHEETS:
            continue
        cc_name, rows = parse_sheet(wb, wbv, code, amount_col, unplaced)
        if not rows:
            continue
        # Income minus cost, in the sheet's own terms: BAS 3xxx is income, and
        # that is the same rule the app infers a kind with (lib/budget/index.ts).
        totals[code] = sum(
            r["_value"] * (1 if r["account"].startswith("3") else -1) for r in rows
        )
        for r in rows:
            out_rows.append({
                "cost_center": code, "cc_name": cc_name,
                "account": r["account"], "account_name": r["account_name"],
                "description": r["description"],
                "quantity": r["quantity"], "unit_price": r["unit_price"],
                "amount": r["amount"],
            })

    fields = ["cost_center", "cc_name", "account", "account_name",
              "description", "quantity", "unit_price", "amount"]
    # UTF-8 with a BOM, because that is what Excel needs to open a CSV as UTF-8
    # without mangling å ä ö, and the importer strips it (lib/domain/csv.ts).
    with open(args.out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(out_rows)

    sheets = len({r["cost_center"] for r in out_rows})
    print(f"Wrote {len(out_rows)} rows across {sheets} cost centres to {args.out}")

    # ── Reconciliation ────────────────────────────────────────────────
    # The point of the exercise: a sheet whose lines do not add up to the
    # Resultat the overview reports is either a converter that lost rows or a
    # spreadsheet summarised by hand. Importing the difference is how a budget
    # comparison ends up quietly wrong.
    stated = rambudget_results(wb, wbv, amount_col)
    bad = []
    for code, got in sorted(totals.items()):
        want = stated.get(code)
        if want is None:
            bad.append(f"    {code:8} no Resultat in Rambudget to check against")
        elif abs(got - want) > 1.0:
            bad.append(f"    {code:8} lines sum to {got:>14,.2f}, Rambudget says {want:>14,.2f}"
                       f"  (off by {got - want:,.2f})")
    if bad:
        print(f"\n{len(bad)} sheet(s) do NOT reconcile, do not import before "
              f"explaining these:", file=sys.stderr)
        print("\n".join(bad), file=sys.stderr)
    else:
        print(f"All {len(totals)} sheets reconcile against Rambudget.")

    if unplaced:
        print(f"\n{len(unplaced)} line(s) under a range header could not be placed "
              f"on an account and were DROPPED:", file=sys.stderr)
        for sheet, rng, desc in unplaced:
            print(f"    {sheet:8} {rng:12} {desc}", file=sys.stderr)

    return 1 if (bad or unplaced) else 0


if __name__ == "__main__":
    sys.exit(main())
